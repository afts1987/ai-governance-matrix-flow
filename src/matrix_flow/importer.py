"""Import an organisation's own matrix (xlsx/csv) into data/matrix.json using a column mapping.
Standard library + openpyxl. Keeps a backup of the previous matrix in data/backups/."""
from __future__ import annotations
import csv, io, json, pathlib, re, shutil
from datetime import datetime

ROOT = pathlib.Path(__file__).resolve().parents[2]
DATA, BACKUPS = ROOT / "data", ROOT / "data" / "backups"

# Our schema: field -> (label, required, default)
FIELDS = [
    ("id", "ID (opcional; gerado se vazio)", False, ""),
    ("theme", "Tema / categoria", False, "Imported"),
    ("requirement", "Requisito ou risco (texto)", True, ""),
    ("ai_act_basis", "Base legal (AI Act ou outra)", False, "—"),
    ("source_status", "Estado da fonte", False, "Internal proposal (to validate)"),
    ("gdpr", "RGPD", False, "—"),
    ("iso27001", "ISO 27001", False, "—"),
    ("iso42001", "ISO 42001", False, "—"),
    ("proposed_control", "Controlo proposto", False, ""),
    ("implemented_control", "Controlo implementado", False, ""),
    ("expected_evidence", "Evidência esperada", False, ""),
    ("owner", "Responsável", False, ""),
    ("implementation_status", "Estado de implementação", False, ""),
    ("review_date", "Data de revisão", False, ""),
    ("notes", "Notas", False, ""),
]
DEFAULTS = {"row_type": "Legal requirement", "nature": "Legal obligation", "application_date": "—",
            "required_role": "Any", "required_risk_class": "Any", "condition": "—", "sector": "All", "applicability": None,
            "iso27701": "—", "standards": "—", "standard_status": "—", "guidance": "—", "sme_variant": "Same as proposed control",
            "sources": [], "verified_on": "", "to_confirm": ""}

def read_table(filename: str, data: bytes) -> tuple[list[str], list[list]]:
    """Return (headers, rows) from an xlsx (first sheet with a header row) or csv."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t") if text.strip() else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    # header = first row with at least 3 non-empty cells
    for i, r in enumerate(rows):
        if sum(1 for c in r if c not in (None, "")) >= 3:
            headers = [str(c).strip() if c not in (None, "") else f"col{j+1}" for j, c in enumerate(r)]
            body = [r for r in rows[i + 1:] if any(c not in (None, "") for c in r)]
            return headers, body
    return [], []

def guess_mapping(headers: list[str]) -> dict[str, str]:
    """Best-effort automatic mapping by header keywords (PT/EN)."""
    keys = {
        "id": ["id", "ref", "código", "codigo", "n.º", "no."],
        "theme": ["tema", "theme", "categoria", "category", "domínio", "domain", "área", "area"],
        "requirement": ["requisito", "requirement", "risco", "risk", "descrição", "description", "obrigação"],
        "ai_act_basis": ["base legal", "legal basis", "artigo", "article", "ai act", "norma", "fonte legal", "referência legal"],
        "source_status": ["estado da fonte", "source status"],
        "gdpr": ["rgpd", "gdpr"],
        "iso27001": ["27001"],
        "iso42001": ["42001"],
        "proposed_control": ["controlo proposto", "proposed control", "controlo", "control", "mitigação", "mitigation", "medida"],
        "implemented_control": ["controlo implementado", "implemented control", "implementado"],
        "expected_evidence": ["evidência", "evidence", "prova"],
        "owner": ["responsável", "owner", "dono", "accountable"],
        "implementation_status": ["estado de implementação", "implementation status", "estado", "status"],
        "review_date": ["revisão", "review"],
        "notes": ["notas", "notes", "observações", "comments"],
    }
    mapping, used = {}, set()
    for field, kws in keys.items():
        for h in headers:
            hl = h.lower()
            if h not in used and any(k in hl for k in kws):
                mapping[field] = h; used.add(h); break
    return mapping

def build_rows(headers: list[str], body: list[list], mapping: dict[str, str]) -> list[dict]:
    idx = {h: i for i, h in enumerate(headers)}
    out = []
    for n, r in enumerate(body, 1):
        def val(field):
            h = mapping.get(field)
            if not h or h not in idx or idx[h] >= len(r): return None
            v = r[idx[h]]
            return None if v in (None, "") else (v.isoformat() if hasattr(v, "isoformat") else str(v).strip())
        row = dict(DEFAULTS)
        for field, _, _, default in FIELDS:
            row[field] = val(field) or default
        if not row["requirement"]:
            continue
        row["id"] = row["id"] or f"IMP-{n:03d}"
        row["id"] = re.sub(r"\s+", "-", row["id"])
        out.append(row)
    return out

def apply(rows: list[dict], label: str) -> pathlib.Path:
    """Replace data/matrix.json after backing up the current one."""
    BACKUPS.mkdir(parents=True, exist_ok=True)
    cur = DATA / "matrix.json"
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    if cur.exists():
        shutil.copy(cur, BACKUPS / f"matrix_{stamp}.json")
    if not (BACKUPS / "matrix_base.json").exists() and cur.exists():
        shutil.copy(cur, BACKUPS / "matrix_base.json")
    cur.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    (DATA / "matrix_meta.json").write_text(json.dumps({"label": label, "imported_at": stamp, "rows": len(rows)}, ensure_ascii=False, indent=2), encoding="utf-8")
    return cur

def restore_base() -> bool:
    base = BACKUPS / "matrix_base.json"
    if not base.exists(): return False
    shutil.copy(base, DATA / "matrix.json")
    meta = DATA / "matrix_meta.json"
    if meta.exists(): meta.unlink()
    return True


# ---------------------------------------------------------------------------------------------------------------
# Round-trip: sync the organisation's own columns (yellow) from an exported workbook back into matrix.json, by ID.
SYNC_FIELDS = {"Implemented control": "implemented_control", "Owner": "owner",
               "Implementation status": "implementation_status", "Review date": "review_date"}

def sync_from_export(filename: str, data: bytes, actor: str = "human:excel-sync") -> dict:
    """Read the 'Matrix' sheet of an exported workbook and update only the yellow columns of rows matched by ID.
    Every change goes through store.update_field, so it lands in history.json with the actor."""
    import sys
    sys.path.insert(0, str(ROOT / "src"))
    from matrix_flow import store
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    from matrix_flow.i18n import HEADERS, VALUES_BACK
    ws = wb["Matrix"] if "Matrix" in wb.sheetnames else (wb["Matriz"] if "Matriz" in wb.sheetnames else wb.worksheets[0])
    pt2en = {HEADERS[k]: k for k in SYNC_FIELDS if k in HEADERS}
    rows = list(ws.iter_rows(values_only=True))
    # header = first row containing 'ID' and at least one sync column
    hdr_i = next((i for i, r in enumerate(rows) if r and "ID" in r and any((h in r) or (HEADERS.get(h) in r) for h in SYNC_FIELDS)), None)
    if hdr_i is None:
        return {"error": "Não encontrei um cabeçalho com 'ID' e as colunas da organização (Implemented control, Owner, Implementation status, Review date)."}
    headers = [pt2en.get(str(c).strip(), str(c).strip()) if c is not None else "" for c in rows[hdr_i]]
    idx = {h: headers.index(h) for h in SYNC_FIELDS if h in headers}
    id_col = headers.index("ID")
    known = {r["id"] for r in store.load("matrix")}
    changed, unknown, seen = 0, [], 0
    for r in rows[hdr_i + 1:]:
        if not r or r[id_col] in (None, ""): continue
        rid = str(r[id_col]).strip(); seen += 1
        if rid not in known: unknown.append(rid); continue
        for header, field in SYNC_FIELDS.items():
            if header not in idx: continue
            v = r[idx[header]]
            v = None if v in (None, "") else (v.isoformat() if hasattr(v, "isoformat") else str(v).strip())
            if isinstance(v, str): v = VALUES_BACK.get(v, v)  # Portuguese list value -> stored English value
            cur = next(x for x in store.load("matrix") if x["id"] == rid).get(field)
            if (cur or None) != v:
                store.update_field("matrix", rid, field, v, actor=actor, basis=f"excel sync: {filename}")
                changed += 1
    return {"rows": seen, "changed": changed, "unknown": unknown[:20]}
