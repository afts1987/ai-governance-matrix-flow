"""Export the current state (data/*.json) as a versioned, formatted Excel workbook.
Called automatically at the end of every persisted round; can also be run by hand:
    python scripts/export_matrix.py            -> exports/AI_Governance_Risk_Matrix_v<ver>_<date>.xlsx
Version: data/version.json holds the base (e.g. 1.2); each export increments the third number (1.2.1, 1.2.2, ...)."""
from __future__ import annotations
import json, pathlib, sys
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
DATA, OUT = ROOT / "data", ROOT / "exports"
import os, sys as _sys0; _sys0.path.insert(0, str(ROOT / "src"))
from matrix_flow.i18n import h as _h, v as _v, sheet as _sheet
LANG = os.getenv("EXPORT_LANG", "pt").lower()

AR = "Arial"; NAVY = "1F3A5F"; TEAL = "2E6E8E"; LIGHT = "F4F6F8"; YEL = "FFF4CC"; GRN = "E6F2EA"
fH = Font(name=AR, bold=True, color="FFFFFF", size=10); fB = Font(name=AR, size=9); fBB = Font(name=AR, size=9, bold=True)
fT = Font(name=AR, size=16, bold=True, color=NAVY); fS = Font(name=AR, size=10, color="5B6B7A"); fH2 = Font(name=AR, size=11, bold=True, color=TEAL)
P = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="D0D7DE"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top"); CT = Alignment(wrap_text=True, vertical="center", horizontal="center")

def load(name, default):
    p = DATA / f"{name}.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default

def next_version() -> str:
    p = DATA / "version.json"
    v = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {"base": "1.2", "export": 0}
    v["export"] = int(v.get("export", 0)) + 1
    p.write_text(json.dumps(v, indent=2), encoding="utf-8")
    return f"{v['base']}.{v['export']}"

def sheet(wb, title, headers, widths, rows, title_text, subtitle, zebra=True, yellow_cols=(), freeze="B4"):
    ws = wb.create_sheet(_sheet(title, LANG))
    headers = [_h(x, LANG) for x in headers]
    rows = [[_v(c, LANG) for c in r] for r in rows]
    ws["A1"] = title_text; ws["A1"].font = fT
    ws["A2"] = subtitle; ws["A2"].font = fS
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        x = ws.cell(3, c, h); x.font = fH; x.fill = P(NAVY); x.alignment = CT; x.border = BD
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[3].height = 40
    for i, r in enumerate(rows):
        for c, v in enumerate(r, 1):
            x = ws.cell(4 + i, c, v); x.font = fB; x.alignment = WR; x.border = BD
            if c in yellow_cols: x.fill = P(YEL)
            elif zebra and i % 2: x.fill = P(LIGHT)
    ws.freeze_panes = freeze; ws.sheet_view.showGridLines = False
    if rows: ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{3 + len(rows)}"
    return ws

T = {
 "Organisation profile": "Perfil da organização", "Answers that determine applicability in the Matrix sheet.": "Respostas que determinam a aplicabilidade na folha Matriz.",
 "Requirement catalogue": "Catálogo de requisitos", "Source register": "Registo de fontes",
 "Status and stage as approved at the human gate; notes carry the date and the evidence found.": "Estado e fase aprovados no gate humano; as notas trazem a data e a evidência encontrada.",
 "Rounds": "Rondas", "One report per persisted round, approved at the second human gate.": "Um relatório por ronda gravada, aprovado no segundo gate humano.",
 "Audit log": "Registo de auditoria", "Every write to the store: human:<gate> for approved decisions, human:auto-reconfirm for silent reconfirmations.": "Cada escrita: human:<gate> para decisões aprovadas, human:auto-reconfirm para reconfirmações silenciosas.",
 "Glossário": "Glossário", "Siglas e termos usados neste ficheiro, na consola e na documentação.": "Siglas e termos usados neste ficheiro, na consola e na documentação.",
}
def tt(x): return T.get(x, x) if LANG == "pt" else x

def main(version: str | None = None) -> pathlib.Path:
    matrix, profile, sources = load("matrix", []), load("profile", {}), load("sources", [])
    history, reports = load("history", []), load("reports", [])
    meta = load("matrix_meta", None)
    ver = version or next_version(); today = date.today().isoformat()
    wb = Workbook(); wb.remove(wb.active)

    # Read me
    lm = wb.create_sheet(_sheet("Read me", LANG)); lm.sheet_view.showGridLines = False
    last = reports[-1] if reports else None
    lines_pt = [
        ("Matriz de Risco de Governação de IA", fT),
        (f"Versão {ver} — exportada a {today} a partir do repositório de dados (data/*.json). {len(matrix)} linhas, {len(sources)} fontes, {len(reports)} rondas, {len(history)} escritas auditadas."
         + (f" Catálogo: {meta['label']} (importado a {meta['imported_at']})." if meta else " Catálogo: matriz de base."), fS),
        ("", fB),
        ("O que é este ficheiro", fH2),
        ("Uma fotografia da matriz tal como é mantida pelo Flow CrewAI: fontes verificadas por agentes e aprovadas por uma pessoa, presença de evidência verificada contra a loja de evidência, cada alteração registada com ator e base. O JSON é a fonte de verdade; este livro é a vista. O texto das linhas (requisitos, controlos, evidências) está em inglês nesta versão.", fB),
        ("", fB),
        ("Última ronda", fH2),
        ((f"{last['round']} — {last['summary']}" if last else "Ainda não há rondas gravadas."), fB),
        ("", fB),
        ("Folhas", fH2),
        ("Perfil · Matriz (aplicabilidade calculada para este perfil) · Fontes (estado, fase, data de verificação, notas aprovadas) · Rondas (um relatório por ronda) · Histórico (cada escrita: quando, o quê, quem, com base em quê) · Glossário (siglas e termos)", fB),
        ("", fB),
        ("Regras impostas em código", fH2),
        ("Os agentes nunca escrevem: classe de risco, aplicabilidade, estado de implementação, controlo implementado, responsável. A descarga de texto limita-se a domínios oficiais. Texto descarregado é dados, não instruções. Uma proposta igual ao registado é reconfirmação, não mudança. Evidência mais antiga do que a última verificação baixa a confiança. A presença de evidência é verificada de forma determinística, sem modelo.", fB),
        ("", fB),
        ("Aviso", fH2),
        ("Ferramenta de trabalho, não parecer jurídico. As notas das fontes registam o que um agente encontrou e uma pessoa aprovou, com o URL de evidência; confirmar na fonte primária antes de confiar nelas. Licença CC BY 4.0.", fB),
    ]
    lines_en = [
        ("AI Governance Risk Matrix", fT),
        (f"Version {ver} — exported {today} from the live data store (data/*.json). {len(matrix)} rows, {len(sources)} sources, {len(reports)} rounds, {len(history)} audited writes."
         + (f" Catalogue: {meta['label']} (imported {meta['imported_at']})." if meta else " Catalogue: base matrix."), fS),
        ("", fB),
        ("What this file is", fH2),
        ("A snapshot of the matrix as maintained by the CrewAI Flow: sources verified by agents and approved by a person, evidence presence checked against the evidence store, every change logged with actor and basis. The JSON store is the source of truth; this workbook is the view.", fB),
        ("", fB),
        ("Latest round", fH2),
        ((f"{last['round']} — {last['summary']}" if last else "No rounds persisted yet."), fB),
        ("", fB),
        ("Sheets", fH2),
        ("Profile · Matrix (applicability as computed for this profile) · Sources (status, stage, verified on, approved notes) · Rounds (one report per round) · History (every write: when, what, who, basis) · Glossary (acronyms and terms)", fB),
        ("", fB),
        ("Rules enforced in code", fH2),
        ("Agents never write: required risk class, applicability, implementation status, implemented control, owner. Fetching is limited to official domains. Fetched text is data, not instructions. A proposal equal to what is on record is a reconfirmation, not a change. Evidence older than the last verification lowers confidence. Evidence presence is checked deterministically; no model is involved.", fB),
        ("", fB),
        ("Disclaimer", fH2),
        ("Working tool, not legal advice. Source notes record what an agent found and a person approved, with the evidence URL; verify against the primary source before relying on them. Licence CC BY 4.0.", fB),
    ]
    lines = lines_pt if LANG == "pt" else lines_en
    for i, (t, f) in enumerate(lines, 1):
        c = lm.cell(i, 1, t); c.font = f; c.alignment = WR
    lm.column_dimensions["A"].width = 130

    # Profile
    sheet(wb, "Profile", ["Key", "Answer"], [26, 40], [[k, v] for k, v in profile.items()],
          "Organisation profile", "Answers that determine applicability in the Matrix sheet.", freeze="A4")

    # Matrix
    cols = ["ID","Theme","Row type","Nature","Requirement / risk","AI Act basis","Source status (AI Act)","Application date",
            "Required role","Required risk class","Additional condition","Sector","Applicability",
            "GDPR","ISO/IEC 27001:2022 Annex A","ISO/IEC 42001:2023 Annex A","ISO/IEC 27701:2025 Annex A","European / other standards","Standard status",
            "Guidance","Proposed control","Proportionate variant (SME)","Expected evidence","Implemented control","Owner","Implementation status","Review date",
            "Sources","Verified on","To confirm","Notes"]
    keys = ["id","theme","row_type","nature","requirement","ai_act_basis","source_status","application_date","required_role","required_risk_class","condition","sector","applicability",
            "gdpr","iso27001","iso42001","iso27701","standards","standard_status","guidance","proposed_control","sme_variant","expected_evidence","implemented_control","owner",
            "implementation_status","review_date","sources","verified_on","to_confirm","notes"]
    widths = [10,16,16,17,52,20,22,24,13,16,15,13,15,22,24,24,28,30,22,26,42,32,34,30,18,16,12,12,12,30,36]
    # Improvement opportunity: derived, never typed by an agent. Proposed control without implementation, or evidence missing in the last round.
    last_ev = {e["row_id"]: e for e in (reports[-1].get("evidence", []) if reports else [])}
    def opportunity(r):
        parts = []
        if r.get("proposed_control") and (r.get("implementation_status") in (None, "", "Not started")):
            parts.append("Implement: " + str(r["proposed_control"])[:160])
        e = last_ev.get(r["id"])
        if e and e.get("found") in ("not_found", "outdated"):
            parts.append(("Evidence missing: " if e["found"] == "not_found" else "Evidence outdated: ") + str(r.get("expected_evidence") or "")[:120])
        return " | ".join(parts) if r.get("applicability") == "Applicable" and parts else ""
    cols.insert(27, "Improvement opportunity (derived)"); keys.insert(27, "__opportunity"); widths.insert(27, 44)
    for r in matrix: r["__opportunity"] = opportunity(r)
    rows = [[("; ".join(r[k]) if isinstance(r.get(k), list) else r.get(k)) for k in keys] for r in matrix]
    ws = sheet(wb, "Matrix", cols, widths, rows, tt("Requirement catalogue"), (f"Aplicabilidade calculada para o perfil acima a {today}. As colunas amarelas são o registo da organização." if LANG == "pt" else f"Applicability computed for the profile above on {today}. Yellow columns are the organisation's own record."),
               yellow_cols=(24, 25, 26, 27), freeze="F4")
    for k, v in {"Final text — in force":"DDEFD6","Final text — deferred application":"DCE9F7","Harmonised standard cited in the OJ":"BFE0B5","Draft (prEN)":"FBE3D4","Non-binding guidance":"EDEDED","Internal proposal (to validate)":"FFF0B3"}.items():
        ws.conditional_formatting.add("G4:G400", FormulaRule(formula=[f'$G4="{_v(k, LANG)}"'], fill=P(v)))
    for k, v in {"Applicable":"CDE8D2","Depends":"FFF0B3","Not applicable":"E6E6E6"}.items():
        ws.conditional_formatting.add("M4:M400", FormulaRule(formula=[f'$M4="{_v(k, LANG)}"'], fill=P(v)))

    # Sources
    srows = [[s.get("id"), s.get("source"), s.get("type"), s.get("status"), s.get("stage"), s.get("reference"), s.get("verified_on"), s.get("how_verified"), s.get("note")] for s in sources]
    ws = sheet(wb, "Sources", ["ID","Source","Type","Status","Stage","Reference","Verified on","How it was verified","Note (agent finding, approved by a person)"],
               [7,44,22,28,16,40,12,36,60], srows, tt("Source register"), tt("Status and stage as approved at the human gate; notes carry the date and the evidence found."))
    for k, v in {"Final text — in force":"DDEFD6","Final text — deferred application":"DCE9F7","Harmonised standard cited in the OJ":"BFE0B5","Draft (prEN)":"FBE3D4","Non-binding guidance":"EDEDED","Published standard (not harmonised)":"DCE9F7"}.items():
        ws.conditional_formatting.add("D4:D400", FormulaRule(formula=[f'$D4="{_v(k, LANG)}"'], fill=P(v)))

    # Rounds
    rrows = [[r.get("round"), r.get("evidence_unknown_count"), len(r.get("evidence_missing", [])), "; ".join(r.get("sources_changed", [])), r.get("summary")] for r in reports]
    sheet(wb, "Rounds", ["Round","Evidence not checked","Evidence missing","Sources changed","Summary"], [22,12,12,60,80], rrows, tt("Rounds"), tt("One report per persisted round, approved at the second human gate."), freeze="B4")

    # History
    hrows = [[h.get("ts"), h.get("entity"), h.get("id"), h.get("field"), str(h.get("old")), str(h.get("new")), h.get("actor"), h.get("basis")] for h in history]
    sheet(wb, "History", ["When (UTC)","Entity","ID","Field","Before","After","Actor","Basis"], [20,10,10,16,40,40,22,50], hrows, tt("Audit log"), tt("Every write to the store: human:<gate> for approved decisions, human:auto-reconfirm for silent reconfirmations."), freeze="B4")

    # Glossary
    import sys as _sys; _sys.path.insert(0, str(ROOT / "src"))
    from matrix_flow.glossary import GLOSSARY
    grows = [[sec, k, v] for sec, items in GLOSSARY for k, v in items]
    sheet(wb, "Glossary", ["Secção", "Sigla / termo", "Significado"], [18, 34, 110], grows, tt("Glossário"), tt("Siglas e termos usados neste ficheiro, na consola e na documentação."), freeze="A4")

    OUT.mkdir(exist_ok=True)
    out = OUT / f"AI_Governance_Risk_Matrix_v{ver}_{today}.xlsx"
    wb.save(out)
    return out

if __name__ == "__main__":
    print(main(sys.argv[1] if len(sys.argv) > 1 else None))
