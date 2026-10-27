"""Convert the AI Governance Risk Matrix workbook into the JSON files the Flow reads.

Usage: python scripts/xlsx_to_json.py AI_Governance_Risk_Matrix_v1.2.xlsx data/
The Excel file stays the human-facing view; JSON is the source the agents work from.
"""
import json, sys, pathlib
from openpyxl import load_workbook

COLS = ["id","theme","row_type","nature","requirement","ai_act_basis","source_status","application_date",
        "required_role","required_risk_class","condition","sector","applicability",
        "gdpr","iso27001","iso42001","iso27701","standards","standard_status","guidance",
        "proposed_control","sme_variant","expected_evidence","implemented_control","owner",
        "implementation_status","review_date","sources","verified_on","to_confirm","notes"]

def main(xlsx, out_dir):
    wb = load_workbook(xlsx, data_only=True)
    out = pathlib.Path(out_dir); out.mkdir(parents=True, exist_ok=True)

    ws = wb["Matrix"]
    rows = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r[0]:
            continue
        d = dict(zip(COLS, r[:len(COLS)]))
        d["applicability"] = None  # recomputed by the Flow from profile.json
        d["sources"] = [s.strip() for s in (d["sources"] or "").split(";") if s.strip()]
        rows.append(d)
    (out / "matrix.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")

    pf = wb["Profile"]
    profile = {}
    for r in pf.iter_rows(min_row=5, values_only=True):
        if r[0]:
            profile[r[0]] = r[2]
    (out / "profile.json").write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")

    fs = wb["Sources"]
    sources = []
    for r in fs.iter_rows(min_row=4, values_only=True):
        if r[0]:
            sources.append(dict(zip(["id","source","type","status","reference","verified_on","how_verified","note"], r[:8])))
    (out / "sources.json").write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(rows)} rows, {len(profile)} profile answers, {len(sources)} sources -> {out}")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "data")
